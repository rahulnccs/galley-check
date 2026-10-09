"""Maintain the journal list from a spreadsheet.

    python scripts/journals_sheet.py export journals.xlsx [--blank]
    python scripts/journals_sheet.py import journals.xlsx

`export` writes the journals in galley/profiles/journals/ to an Excel
workbook, one row per journal (with --blank, just the headers and an
example row: the template). `import` writes each row back as
galley/profiles/journals/<id>.json, replacing any file with that id.
Every row is checked first, and nothing is written unless all are valid.
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import date, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from galley.checks.offline.submission import JOURNAL_DIR, Profile  # noqa: E402

LIMITS = ["title_words", "title_chars", "abstract_words", "main_text_words",
          "total_words", "keywords", "references", "figures", "tables",
          "display_items"]
# header, help shown on the header cell, required
COLUMNS = [
    ("id", "File name: lower case with hyphens. Never change it once published.", True),
    ("name", "The journal's name, e.g. Nature Communications.", True),
    ("publisher", "e.g. Springer Nature, Cell Press, PLOS.", False),
    ("article_type", "e.g. Article, Research Article, Original Research.", False),
    ("guidelines_url", "The journal's author guidelines page.", True),
    ("verified", "The date you read the guidelines (YYYY-MM-DD). Blank = not yet checked.",
     False),
    ("title_words", "Longest title, in words.", False),
    ("title_chars", "Longest title, in characters including spaces.", False),
    ("abstract_words", "Longest abstract, in words.", False),
    ("abstract_headings", "Structured abstract headings, separated by ; "
                          "e.g. Background; Results; Conclusions", False),
    ("main_text_words", "Introduction to discussion, in words (methods not counted).",
     False),
    ("total_words", "Abstract, main text and methods, in words.", False),
    ("keywords", "Most keywords allowed.", False),
    ("references", "Most references in the list.", False),
    ("figures", "Most main figures.", False),
    ("tables", "Most main tables.", False),
    ("display_items", "Most figures and tables together.", False),
    ("reference_style", "numbered or author_year", False),
    ("max_authors_listed", "Authors listed in a reference before et al.", False),
    ("require_doi", "Yes if every reference needs a DOI.", False),
    ("required_sections", "Sections or statements that must be present, separated by ; "
                          "e.g. Data availability; Competing interests", False),
    ("notes", "Shown to users.", False),
]
EXAMPLE = {"id": "example-journal", "name": "Example Journal",
           "publisher": "Example Publisher", "article_type": "Research Article",
           "guidelines_url": "https://example.org/authors", "verified": "2026-10-01",
           "abstract_words": 250, "main_text_words": 5000, "keywords": 6,
           "abstract_headings": "Background; Methods; Results; Conclusions",
           "reference_style": "numbered", "required_sections": "Data availability",
           "notes": "Invented example: rows whose id starts with example- are skipped."}


def row_of(data: dict) -> dict:
    r = {k: data.get(k) for k in ("id", "name", "publisher", "article_type",
                                  "guidelines_url", "verified", "reference_style",
                                  "max_authors_listed", "notes")}
    r.update({k: (data.get("limits") or {}).get(k) for k in LIMITS})
    r["abstract_headings"] = "; ".join(data.get("abstract_headings") or [])
    r["required_sections"] = "; ".join(data.get("required_sections") or [])
    r["require_doi"] = "Yes" if data.get("require_doi") else None
    return r


def export(path: Path, blank: bool = False) -> int:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    book = Workbook()
    sheet = book.active
    sheet.title = "Journals"
    sheet.append([c[0] for c in COLUMNS])
    for i, (key, help_text, required) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(1, i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="C0504D" if required else "4A4E54")
        cell.comment = Comment(help_text, "Galley")
        sheet.column_dimensions[cell.column_letter].width = (
            44 if key in ("guidelines_url", "required_sections", "notes", "name") else 16)
    rows = [EXAMPLE]
    if not blank:
        rows += [row_of(json.loads(p.read_text(encoding="utf-8")))
                 for p in sorted(JOURNAL_DIR.glob("*.json"))]
    grey = Font(color="8E8E93", italic=True)
    for n, r in enumerate(rows):
        sheet.append([r.get(c[0]) for c in COLUMNS])
        if n == 0:
            for cell in sheet[sheet.max_row]:
                cell.font = grey
    sheet.freeze_panes = "C2"
    style = DataValidation(type="list", formula1='"numbered,author_year"', allow_blank=True)
    yes = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    sheet.add_data_validation(style)
    sheet.add_data_validation(yes)
    col = {c[0]: i for i, c in enumerate(COLUMNS, start=1)}
    letter = lambda k: sheet.cell(1, col[k]).column_letter   # noqa: E731
    style.add(f"{letter('reference_style')}2:{letter('reference_style')}2000")
    yes.add(f"{letter('require_doi')}2:{letter('require_doi')}2000")
    book.save(path)
    return len(rows) - 1


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def entry(r: dict) -> dict:
    """One row as a journal file's contents."""
    verified = r.get("verified")
    if isinstance(verified, datetime):
        verified = verified.date()
    if isinstance(verified, date):
        verified = verified.isoformat()
    e = {"id": _text(r.get("id")), "name": _text(r.get("name")),
         "publisher": _text(r.get("publisher")) or None,
         "article_type": _text(r.get("article_type")) or None,
         "guidelines_url": _text(r.get("guidelines_url")),
         "verified": _text(verified) or None}
    limits = {}
    for k in LIMITS:
        v = r.get(k)
        if v not in (None, ""):
            try:
                limits[k] = int(float(v))
            except ValueError:
                raise ValueError(f"{k} should be a whole number, not {v!r}") from None
    if limits:
        e["limits"] = limits
    for k in ("abstract_headings", "required_sections"):
        items = [x.strip() for x in _text(r.get(k)).split(";") if x.strip()]
        if items:
            e[k] = items
    if _text(r.get("reference_style")):
        e["reference_style"] = _text(r.get("reference_style")).lower()
    if _text(r.get("max_authors_listed")):
        e["max_authors_listed"] = int(float(r["max_authors_listed"]))
    if _text(r.get("require_doi")).lower() in ("yes", "y", "true", "1"):
        e["require_doi"] = True
    if _text(r.get("notes")):
        e["notes"] = _text(r.get("notes"))
    return {k: v for k, v in e.items() if v is not None or k == "verified"}


def convert(path: Path) -> list[dict]:
    from openpyxl import load_workbook
    sheet = load_workbook(path, data_only=True)["Journals"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(h or "").replace("*", "").strip() for h in next(rows)]
    out, problems, seen = [], [], set()
    for n, values in enumerate(rows, start=2):
        r = dict(zip(headers, values))
        jid = _text(r.get("id"))
        if not any(v not in (None, "") for v in values) or jid.startswith("example-"):
            continue
        try:
            for key in ("id", "name", "guidelines_url"):
                if not _text(r.get(key)):
                    raise ValueError(f"{key} is required")
            if jid in seen:
                raise ValueError("this id is used twice")
            seen.add(jid)
            e = entry(r)
            Profile.from_dict(e)                   # the app's own checks
            out.append(e)
        except (ValueError, TypeError) as err:
            problems.append(f"Row {n} ({jid or 'no id'}): {err}")
    if problems:
        raise ValueError("\n".join(problems))
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("action", choices=["export", "import"])
    ap.add_argument("spreadsheet", type=Path)
    ap.add_argument("--blank", action="store_true", help="export only the template")
    ap.add_argument("--out", type=Path, default=JOURNAL_DIR,
                    help="folder to import into (default: the bundled list)")
    args = ap.parse_args(argv)
    if args.action == "export":
        n = export(args.spreadsheet, args.blank)
        print(f"Wrote {n} journals to {args.spreadsheet}")
        return 0
    try:
        entries = convert(args.spreadsheet)
    except ValueError as err:
        print(f"Nothing written. Fix these first:\n{err}", file=sys.stderr)
        return 1
    args.out.mkdir(parents=True, exist_ok=True)
    for e in entries:
        (args.out / f"{e['id']}.json").write_text(
            json.dumps(e, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    checked = sum(1 for e in entries if e.get("verified"))
    print(f"Wrote {len(entries)} journals ({checked} checked) to {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
