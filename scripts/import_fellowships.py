"""Turn a filled-in fellowship spreadsheet into Galley's database files.

    python scripts/import_fellowships.py filled-template.xlsx

The spreadsheet is templates/fellowship-database-template.xlsx: one row per
fellowship on the Fellowships sheet, one row per required document on the
Documents sheet. Columns are found by their header, so older copies of the
template (without category, career_levels and so on) still work. Rows whose
id starts with "example-" are skipped.

Each fellowship is written to galley/fellowships/data/<id>.json, replacing
any file with that id. Nothing is written unless every row is valid.
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import date, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from galley.fellowships.model import Fellowship  # noqa: E402

DATA_DIR = ROOT / "galley" / "fellowships" / "data"
DEADLINE_COLUMNS = [("call_opens", "call_opens"), ("internal_deadline", "internal"),
                    ("pre_proposal_deadline", "pre_proposal"),
                    ("final_deadline", "final")]


def _rows(sheet) -> list[dict]:
    """The sheet's rows as dicts keyed by header (without the ' *')."""
    rows = sheet.iter_rows(values_only=True)
    headers = [str(h or "").replace("*", "").strip() for h in next(rows)]
    out = []
    for n, values in enumerate(rows, start=2):
        row = {h: _clean(v) for h, v in zip(headers, values) if h}
        if any(v is not None for v in row.values()):
            row["_row"] = n
            out.append(row)
    return out


def _clean(v):
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def _iso(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return (v.date() if isinstance(v, datetime) else v).isoformat()
    return str(v)


def _yes(v) -> bool | None:
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    s = str(v).lower()
    if s in ("yes", "y", "true", "1"):
        return True
    if s in ("no", "n", "false", "0"):
        return False
    raise ValueError(f"expected Yes or No, not {v!r}")


def _list(v, sep=",") -> list[str] | None:
    if v is None:
        return None
    return [x.strip() for x in str(v).split(sep) if x.strip()]


def _number(v):
    if v is None:
        return None
    f = float(v)
    return int(f) if f.is_integer() else f


def _time(v) -> str | None:
    if v is None:
        return None
    if hasattr(v, "strftime"):
        return v.strftime("%H:%M")
    return str(v)


def _document(d: dict) -> dict:
    doc = {"name": d.get("document_name"),
           "max_pages": _number(d.get("max_pages")),
           "max_words": _number(d.get("max_words")),
           "min_font_size": _number(d.get("min_font_size")),
           "min_margin_cm": _number(d.get("min_margin_cm")),
           "file_format": d.get("file_format"),
           "sections": _list(d.get("required_sections"), ";"),
           "template_url": d.get("template_url"),
           "notes": d.get("notes")}
    return {k: v for k, v in doc.items() if v not in (None, [])}


def entry(r: dict, documents: list[dict]) -> dict:
    """One spreadsheet row (and its documents) as a database entry."""
    e = {"id": r.get("id"), "name": r.get("name"), "funder": r.get("funder"),
         "url": r.get("official_url"), "verified": _iso(r.get("verified"))}
    for key in ("category", "purpose", "membership"):
        if r.get(key):
            e[key] = r[key]
    if r.get("career_levels"):
        e["career_levels"] = _list(r["career_levels"])
    if r.get("membership_min_months"):
        e["membership_min_months"] = _number(r["membership_min_months"])
    e["fields"] = _list(r.get("fields")) or ["any"]
    if r.get("track") and r["track"] != "any":
        e["track"] = r["track"]

    stage = {"phd_required": _yes(r.get("phd_required")),
             "min_years_since_phd": _number(r.get("min_years_since_phd")),
             "max_years_since_phd": _number(r.get("max_years_since_phd")),
             "career_breaks_extend": _yes(r.get("career_breaks_extend"))}
    stage = {k: v for k, v in stage.items() if v is not None}
    if stage:
        e["career_stage"] = stage

    for key in ("nationalities", "excluded_nationalities", "residence",
                "host_countries", "excluded_host_countries"):
        if r.get(key):
            e[key] = [c.upper() for c in _list(r[key])]
    months, years = r.get("mobility_max_months_in_host"), r.get("mobility_window_years")
    if months is not None and years is not None:
        e["mobility"] = {"max_months_in_host": _number(months),
                         "window_years": _number(years)}
    elif months is not None or years is not None:
        raise ValueError("fill in both mobility columns, or neither")

    estimated = bool(_yes(r.get("deadline_estimated")))
    deadlines = []
    for column, kind in DEADLINE_COLUMNS:
        if r.get(column) is None:
            continue
        d = {"kind": kind, "date": _iso(r[column])}
        if kind == "final":
            if r.get("deadline_time"):
                d["time"] = _time(r["deadline_time"])
            if r.get("timezone"):
                d["timezone"] = r["timezone"]
        if estimated:
            d["estimated"] = True
        deadlines.append(d)
    if deadlines:
        e["deadlines"] = deadlines
    for key in ("rolling", "annual"):
        v = _yes(r.get(key))
        if v is not None and v != (key == "annual"):    # leave out defaults
            e[key] = v

    if r.get("amount") is not None:
        e["amount"] = str(r["amount"])
    if r.get("duration_months") is not None:
        e["duration_months"] = _number(r["duration_months"])
    if r.get("other_rules"):
        e["other_rules"] = _list(r["other_rules"], ";")
    if r.get("notes"):
        e["notes"] = r["notes"]

    req = {"documents": [_document(d) for d in documents],
           "cv_format": r.get("cv_format"), "cv_notes": r.get("cv_notes"),
           "host_letter": _yes(r.get("host_letter")),
           "submission": r.get("submission"),
           "submission_url": r.get("submission_url")}
    req = {k: v for k, v in req.items() if v not in (None, [], False)}
    if req:
        e["requirements"] = req
    return e


def convert(path: Path) -> list[dict]:
    """Every real entry in the spreadsheet, checked; raises on any problem."""
    from openpyxl import load_workbook
    book = load_workbook(path, data_only=True)
    rows = [r for r in _rows(book["Fellowships"])
            if not str(r.get("id") or "").startswith("example-")]
    docs: dict[str, list[dict]] = {}
    if "Documents" in book.sheetnames:
        for d in _rows(book["Documents"]):
            fid = d.get("fellowship_id")
            if not str(fid or "").startswith("example-"):
                docs.setdefault(fid, []).append(d)

    problems, entries, seen = [], [], set()
    for r in rows:
        where = f"Fellowships row {r['_row']} ({r.get('id') or 'no id'})"
        try:
            for key in ("id", "name", "funder", "official_url", "verified"):
                if r.get(key) is None:
                    raise ValueError(f"{key} is required")
            if r["id"] in seen:
                raise ValueError("this id is used twice")
            seen.add(r["id"])
            e = entry(r, docs.get(r["id"], []))
            Fellowship.from_dict(e)                    # the app's own checks
            entries.append(e)
        except ValueError as err:
            problems.append(f"{where}: {err}")
    for fid, ds in docs.items():
        if fid not in seen:
            problems.append(f"Documents row {ds[0]['_row']}: no fellowship "
                            f"with id {fid!r}")
    if problems:
        raise ValueError("\n".join(problems))
    return entries


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("spreadsheet", type=Path)
    ap.add_argument("--out", type=Path, default=DATA_DIR,
                    help="folder to write to (default: the bundled database)")
    args = ap.parse_args(argv)
    try:
        entries = convert(args.spreadsheet)
    except ValueError as err:
        print(f"Nothing written. Fix these first:\n{err}", file=sys.stderr)
        return 1
    args.out.mkdir(parents=True, exist_ok=True)
    new = 0
    for e in entries:
        target = args.out / f"{e['id']}.json"
        new += not target.exists()
        target.write_text(json.dumps(e, indent=2, ensure_ascii=False) + "\n",
                          encoding="utf-8")
    print(f"Wrote {len(entries)} entries ({new} new, {len(entries) - new} "
          f"updated) to {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
