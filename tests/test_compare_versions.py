"""Comparing an old and a revised version of a manuscript."""
from pathlib import Path

import docx
import pytest

from galley.engine import load
from galley.report.compare_versions import (
    compare, default_output_path, highlight, split_sentences, summarize)


def write(path, paragraphs):
    d = docx.Document()
    d.add_heading("A fictional study", 0)
    d.add_heading("Results", 1)
    for text in paragraphs:
        d.add_paragraph(text)
    d.save(path)
    return load(path)


OLD = ["Worms were raised on compost extracts. Diversity fell with age.",
       "Cytokine levels were similar between groups. The effect was modest."]


def test_unchanged_document_has_no_changes(tmp_path):
    old = write(tmp_path / "old.docx", OLD)
    new = write(tmp_path / "new.docx", OLD)
    assert compare(old, new) == []


def test_added_sentence(tmp_path):
    old = write(tmp_path / "old.docx", OLD)
    new = write(tmp_path / "new.docx", [
        OLD[0] + " We repeated the experiment in a second cohort.", OLD[1]])
    changes = compare(old, new)
    assert [c.kind for c in changes] == ["added"]
    assert "second cohort" in changes[0].sentence


def test_edited_sentence_is_not_reported_as_new(tmp_path):
    old = write(tmp_path / "old.docx", OLD)
    new = write(tmp_path / "new.docx", [
        "Worms were raised on bacterial extracts. Diversity fell with age.",
        OLD[1]])
    changes = compare(old, new)
    assert [c.kind for c in changes] == ["edited"]
    assert changes[0].similarity > 0.55


def test_reordered_sentences_are_not_changes(tmp_path):
    old = write(tmp_path / "old.docx", OLD)
    new = write(tmp_path / "new.docx", [
        "Diversity fell with age. Worms were raised on compost extracts.",
        OLD[1]])
    assert compare(old, new) == []


def test_deletions_are_not_reported(tmp_path):
    """Only the revised document is annotated, so removals aren't highlighted."""
    old = write(tmp_path / "old.docx", OLD)
    new = write(tmp_path / "new.docx", [OLD[0]])
    assert compare(old, new) == []


def test_highlighting_writes_a_copy(tmp_path):
    old = write(tmp_path / "old.docx", OLD)
    new_path = tmp_path / "new.docx"
    new = write(new_path, [OLD[0] + " A brand new sentence appears here.", OLD[1]])
    before = new_path.read_bytes()
    out = highlight(new, compare(old, new), tmp_path / "out.docx")
    assert out.exists()
    assert new_path.read_bytes() == before


def test_highlighted_text_is_the_changed_sentence(tmp_path):
    old = write(tmp_path / "old.docx", OLD)
    new = write(tmp_path / "new.docx",
                [OLD[0] + " A brand new sentence appears here.", OLD[1]])
    out = highlight(new, compare(old, new), tmp_path / "out.docx")
    highlighted = []
    for p in docx.Document(str(out)).paragraphs:
        for r in p.runs:
            if r.font.highlight_color is not None:
                highlighted.append(r.text)
    assert any("brand new sentence" in t for t in highlighted)
    assert not any("Diversity fell with age" in t for t in highlighted)


def test_refuses_to_overwrite_the_revised_file(tmp_path):
    old = write(tmp_path / "old.docx", OLD)
    new = write(tmp_path / "new.docx", OLD)
    with pytest.raises(ValueError):
        highlight(new, [], Path(new.path))


def test_summary_counts(tmp_path):
    old = write(tmp_path / "old.docx", OLD)
    new = write(tmp_path / "new.docx", [OLD[0] + " One more sentence here.", OLD[1]])
    text = summarize(old, new, compare(old, new))
    assert "1 added" in text and "0 edited" in text


@pytest.mark.parametrize("text, count", [
    ("One sentence only.", 1),
    ("First one. Second one.", 2),
    ("Values differed (e.g. p < 0.05). Then we repeated it.", 2),
    ("See Fig. 2 for details. The effect held.", 2),
    ("Smith et al. reported this. We confirmed it.", 2),
])
def test_sentence_splitting_handles_abbreviations(text, count):
    assert len(split_sentences(text)) == count


def test_default_output_path():
    assert default_output_path("/tmp/paper_v2.docx").name == "paper_v2_changes.docx"
