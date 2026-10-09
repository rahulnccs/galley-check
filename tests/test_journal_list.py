"""The journal list: pre-filled journals, the new rules, where problems are
in the text, the spreadsheet and the update download."""
import importlib.util
import json
from pathlib import Path

import pytest

from galley.checks.offline import submission
from galley.checks.offline.submission import (JOURNAL_DIR, Profile, check_submission,
                                              count, where_limit_is_crossed)
from galley.model.document import Document, Paragraph

ROOT = Path(__file__).resolve().parent.parent


def make_doc(sections, title="A study of microbial communities in model hosts"):
    paras = [Paragraph(index=0, text=title, style="Title", section="front_matter")]
    for section, text in sections:
        paras.append(Paragraph(index=len(paras), text=text, style="Normal",
                               section=section))
    return Document(path="t.docx", paragraphs=paras)


def test_every_listed_journal_loads_and_is_marked_unchecked_until_verified():
    files = sorted(JOURNAL_DIR.glob("*.json"))
    assert len(files) >= 40
    for path in files:
        data = json.loads(path.read_text(encoding="utf-8"))
        p = Profile.from_dict(data, path)
        assert data["id"] == path.stem and p.guidelines_url.startswith("https://")
        if not p.verified:
            assert "not yet checked" in (p.notes or "")


def test_listed_journals_are_offered(monkeypatch, tmp_path):
    monkeypatch.setattr(submission, "user_profile_dir", lambda: tmp_path / "profiles")
    names = [p.label for p in submission.available_profiles()]
    assert "eLife · Research Article" in names and "Microbiome · Research" in names
    # A corrected copy saved under the same file name replaces the listed one.
    submission.save_profile({"name": "eLife", "verified": "2026-10-09",
                             "limits": {"abstract_words": 160}}, "elife")
    elife = [p for p in submission.available_profiles() if p.name == "eLife"]
    assert len(elife) == 1 and elife[0].limits == {"abstract_words": 160}


def test_the_sentence_where_a_word_limit_is_crossed_is_pointed_out():
    abstract = ("Short opening sentence here. " * 2 +
                "This sentence crosses the limit right about now. Then more text.")
    doc = make_doc([("abstract", abstract)])
    para, sentence = where_limit_is_crossed(doc, "abstract_words", 10)
    assert para == 1 and sentence == "This sentence crosses the limit right about now."
    issues = check_submission(doc, Profile(name="J", verified="2026-10-01",
                                           limits={"abstract_words": 10}))
    over = next(i for i in issues if "over the J limit" in i.message)
    assert over.para_index == 1 and over.anchor == sentence
    assert "highlighted sentence" in over.suggestion


def test_title_characters_keywords_and_structured_abstract():
    doc = make_doc([("abstract", "Background: why. Methods: how. Results: what."),
                    ("abstract", "Keywords: microbiome; mice; diet; bile acids"),
                    ("introduction", "word " * 20)])
    c = count(doc)
    assert c.title_chars == 47 and c.keywords == 4 and c.keywords_index == 2
    p = Profile(name="J", verified="2026-10-01",
                limits={"title_chars": 40, "keywords": 3},
                abstract_headings=["Background", "Methods", "Results", "Conclusions"])
    issues = check_submission(doc, p)
    text = " | ".join(i.message for i in issues)
    assert "title is 47 characters long, over the J limit of 40 characters" in text
    assert "4 keywords, over the J limit of 3" in text
    headings = next(i for i in issues if "structured abstract" in i.message)
    assert headings.severity == "error" and "Conclusions is missing" in headings.message
    title = next(i for i in issues if "characters long" in i.message)
    assert title.para_index == 0


def test_unchecked_requirements_are_softer_and_say_so():
    doc = make_doc([("abstract", "word " * 30), ("introduction", "word " * 10)])
    p = Profile(name="J", limits={"abstract_words": 20},
                required_sections=["Data availability"])
    issues = check_submission(doc, p)
    over = next(i for i in issues if "over the J limit" in i.message)
    assert "(not yet confirmed)" in over.message
    missing = next(i for i in issues if "Data availability" in i.message)
    assert missing.severity == "warning" and "(not yet confirmed)" in missing.message
    assert any("haven't been checked" in i.message for i in issues)


def test_missing_keywords_line():
    doc = make_doc([("abstract", "word " * 10)])
    issues = check_submission(doc, Profile(name="J", verified="2026-10-01",
                                           limits={"keywords": 6}))
    assert any("no “Keywords:” line was found" in i.message for i in issues)


def test_spreadsheet_round_trip(tmp_path):
    pytest.importorskip("openpyxl")
    spec = importlib.util.spec_from_file_location(
        "journals_sheet", ROOT / "scripts" / "journals_sheet.py")
    sheet = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(sheet)
    book = tmp_path / "journals.xlsx"
    assert sheet.main(["export", str(book)]) == 0
    out = tmp_path / "out"
    assert sheet.main(["import", str(book), "--out", str(out)]) == 0
    for path in JOURNAL_DIR.glob("*.json"):
        assert json.loads((out / path.name).read_text()) == json.loads(path.read_text())

    from openpyxl import load_workbook
    wb = load_workbook(book)
    ws = wb["Journals"]
    ws.append(["new-journal", "New Journal", None, None, None])      # no guidelines URL
    wb.save(book)
    assert sheet.main(["import", str(book), "--out", str(tmp_path / "x")]) == 1
    assert not (tmp_path / "x").exists()


def test_journal_list_download(tmp_path):
    good = json.dumps({"id": "j", "name": "J", "guidelines_url": "https://j.org",
                       "verified": None, "limits": {"abstract_words": 200}}).encode()
    bad = json.dumps({"name": "Bad", "reference_style": "footnotes"}).encode()
    files = {"j.json": good, "bad.json": bad}

    class Reply:
        def __init__(self, body):
            self.body = body

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def read(self):
            return self.body

    def opener(req, timeout=0):
        url = req.full_url
        if "contents/journals" in url:
            return Reply(json.dumps([{"name": n, "type": "file",
                                      "download_url": "https://raw/" + n}
                                     for n in files]).encode())
        return Reply(files[url.rsplit("/", 1)[1]])

    folder = tmp_path / "journal-list"
    result = submission.fetch_journal_updates(folder, opener)
    assert result.count == 1 and result.skipped == ["bad.json"]
    assert (folder / "j.json").exists() and submission.journals_last_updated(folder)


def test_journal_picker_searches_and_chooses(monkeypatch, tmp_path):
    pytest.importorskip("PySide6")
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QApplication
    app = QApplication.instance() or QApplication([])  # noqa: F841
    monkeypatch.setattr(submission, "user_profile_dir", lambda: tmp_path / "profiles")
    from galley.gui import journal_picker
    monkeypatch.setattr(journal_picker, "journals_last_updated", lambda: None)
    picker = journal_picker.JournalPicker()
    rows = [picker.list.item(i).text() for i in range(picker.list.count())]
    assert rows[0].startswith("No journal limits")
    assert any(r.startswith("Cell Host & Microbe · Article") for r in rows)
    assert "requirements not yet confirmed" in rows[2]
    picker.search.setText("cell press")
    found = [picker.list.item(i).data(Qt.UserRole) for i in range(picker.list.count())]
    found = [p for p in found if p is not None]
    assert found and all(p.publisher == "Cell Press" for p in found)
    picker.accept_choice()
    assert picker.chosen.publisher == "Cell Press"
    picker.search.setText("no such journal at all")
    assert "No journal matches" in picker.list.item(0).text()
