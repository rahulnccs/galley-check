"""The Fellowships screen. Needs Qt; skipped where PySide6 isn't installed."""
from datetime import date

import pytest

pytest.importorskip("PySide6")

from PySide6.QtWidgets import QApplication, QLabel, QPushButton  # noqa: E402

from galley.checks.offline import submission  # noqa: E402
from galley.fellowships import load_applications  # noqa: E402
from galley.fellowships.model import (Researcher, Stay,  # noqa: E402
                                      load_researcher, save_researcher)

from galley.gui.fellowships_page import APPLICATIONS, MATCHES, PROFILE  # noqa: E402

TODAY = date(2026, 10, 1)


@pytest.fixture(scope="module")
def app():
    existing = QApplication.instance()
    if existing is not None:
        return existing
    try:
        return QApplication([])
    except Exception as e:
        pytest.skip(f"Qt cannot start here: {e}")


@pytest.fixture
def settings(tmp_path, monkeypatch):
    """Point Galley's settings folder at a temporary one."""
    monkeypatch.setattr(submission, "user_profile_dir", lambda: tmp_path / "profiles")
    # Include the bundled example entries so there is something to match.
    import galley.fellowships.model as model
    real = model.load_fellowships

    def with_examples(*a, **k):
        k["include_templates"] = True
        return real(*a, **k)
    import galley.gui.fellowships_page as page_mod
    monkeypatch.setattr(page_mod, "load_fellowships", with_examples)
    return tmp_path


def texts(widget):
    return " | ".join(l.text() for l in widget.findChildren(QLabel))


def make_page(app):
    from galley.gui.fellowships_page import FellowshipsPage
    return FellowshipsPage(today=TODAY)


def test_empty_profile_prompts_to_fill_it_in(app, settings):
    page = make_page(app)
    assert "Fill in your profile" in texts(page.tab_stack.widget(MATCHES))


def test_matches_grouped_by_eligibility(app, settings):
    save_researcher(Researcher(phd_date=date(2025, 1, 1), nationalities=["IN"],
                               residence="IN", clinical=False,
                               target_hosts=["DE"],
                               stays=[Stay("IN", date(2015, 1, 1))]))
    page = make_page(app)
    shown = texts(page.tab_stack.widget(MATCHES))
    assert "YOU'RE ELIGIBLE" in shown
    assert "Example International Postdoctoral Fellowship" in shown
    assert "Fill in your profile" not in shown


def test_add_application_and_tick_a_step(app, settings):
    page = make_page(app)
    page.open_fellowship("example-international-postdoc")
    detail = texts(page.detail)
    assert "WHAT YOU NEED" in detail and "Research proposal" in detail
    page._add_application("example-international-postdoc")
    saved = load_applications(settings / "applications.json")
    assert "example-international-postdoc" in saved
    app_obj = page.apps["example-international-postdoc"]
    from PySide6.QtWidgets import QPushButton
    ticks = [b for b in page.detail.findChildren(QPushButton) if b.isCheckable()]
    ticks[0].setChecked(True)
    assert app_obj.done
    assert load_applications(settings / "applications.json")[
        "example-international-postdoc"].done == app_obj.done


def test_reminders_counted_for_the_banner(app, settings):
    page = make_page(app)
    page._add_application("example-international-postdoc")
    page.today = date(2026, 12, 20)      # a few weeks before the deadline
    assert page.attention_count() > 0
    page.refresh(APPLICATIONS)
    assert "COMING UP" in texts(page.tab_stack.widget(APPLICATIONS))


def test_profile_form_saves(app, settings):
    page = make_page(app)
    page.show_tab(PROFILE)
    page.p_phd_state.setCurrentIndex(0)
    page.p_nationality.setText("in, gb")
    page.p_residence.setText("in")
    page.p_fields["neuroscience"].setChecked(True)
    page.save_profile()
    r = load_researcher(settings / "researcher.json")
    assert r.nationalities == ["IN", "GB"] and r.residence == "IN"
    assert r.fields == ["neuroscience"] and r.phd_date is not None


def test_bad_country_code_is_not_saved(app, settings, monkeypatch):
    from PySide6.QtWidgets import QMessageBox
    warned = []
    monkeypatch.setattr(QMessageBox, "warning", lambda *a: warned.append(a[2]))
    page = make_page(app)
    page.p_nationality.setText("India")
    page.save_profile()
    assert warned and "INDIA" in warned[0]
    assert not (settings / "researcher.json").exists()


def test_main_window_has_both_sections(app, settings):
    from galley.gui.app import MainWindow
    win = MainWindow("Georgia")
    assert win.switcher.buttons[1].text().startswith("Fellowships")
    win.switcher.select(1)
    assert win.sections.currentWidget() is win.fellowships


def test_attach_and_check_application(app, settings, tmp_path, monkeypatch):
    import docx
    proposal = tmp_path / "proposal.docx"
    d = docx.Document()
    for heading in ("Background", "Aims", "Approach", "Feasibility"):
        d.add_heading(heading, level=1)
        d.add_paragraph("Some text.")
    d.save(proposal)

    page = make_page(app)
    page._add_application("example-international-postdoc")
    assert "CHECK YOUR APPLICATION" in texts(page.detail)
    monkeypatch.setattr(page, "_choose_file", lambda title: str(proposal))
    page._attach("example-international-postdoc", "Research proposal")
    saved = load_applications(settings / "applications.json")
    assert saved["example-international-postdoc"].files == {
        "Research proposal": str(proposal)}

    page._check_application("example-international-postdoc")
    report = texts(page.detail)
    assert "to fix" in report                       # career statement missing
    assert "Not attached yet." in report
    assert 'Includes "Feasibility".' in report


def test_category_filter_and_career_stage(app, settings):
    page = make_page(app)
    page._filter_category(3)                     # Travel
    assert "Example International Postdoctoral" not in texts(page.tab_stack.widget(MATCHES))
    page._filter_category(0)                     # All
    assert "Example International Postdoctoral" in texts(page.tab_stack.widget(MATCHES))
    page.show_tab(PROFILE)
    page.p_level.setCurrentIndex(2)              # PhD student
    page.save_profile()
    assert load_researcher(settings / "researcher.json").career_level == "phd_student"


def test_examples_can_be_shown_to_try_galley(app, tmp_path, monkeypatch):
    # The real loader: examples are hidden until the user asks for them.
    monkeypatch.setattr(submission, "user_profile_dir", lambda: tmp_path / "profiles")
    # A database with only the invented examples, as before real entries.
    import shutil
    import galley.fellowships.model as model
    only_examples = tmp_path / "data"
    only_examples.mkdir()
    for p in model.DATA_DIR.glob("example-*.json"):
        shutil.copy(p, only_examples)
    monkeypatch.setattr(model, "DATA_DIR", only_examples)
    from galley.gui.fellowships_page import FellowshipsPage
    page = FellowshipsPage(today=TODAY)
    assert "Example International" not in texts(page.tab_stack.widget(MATCHES))
    assert "Show invented example entries" in texts(page.tab_stack.widget(MATCHES))
    page._toggle_examples(True)
    shown = texts(page.tab_stack.widget(MATCHES))
    assert "Example International Postdoctoral Fellowship" in shown
    assert "Postdoc fellowship (example)" in shown


def test_tabs_are_profile_matches_applications(app, settings):
    page = make_page(app)
    assert [b.text() for b in page.tabs.buttons][:2] == ["Profile", "Matches"]
    assert page.tabs.index == PROFILE                # empty profile: start there
    save_researcher(Researcher(career_level="postdoc"))
    assert make_page(app).tabs.index == MATCHES      # then open on Matches


def test_stage_picker_has_a_colour_per_stage(app, settings):
    from galley.gui.fellowships_page import STAGE_COLOR, Picker
    page = make_page(app)
    page._add_application("example-international-postdoc")
    picker = [p for p in page.detail.findChildren(Picker)][0]
    assert picker.count() == 7
    assert len({c for _, _, c in picker.items}) == 7  # all different
    picker.setCurrentIndex(1)                         # Submitted
    assert page.apps["example-international-postdoc"].status == "submitted"
    assert STAGE_COLOR["submitted"] in picker.styleSheet()


def test_picker_menu_opens_with_big_rows(app, settings):
    from PySide6.QtWidgets import QPushButton
    from galley.gui.fellowships_page import _PickerMenu
    menu = _PickerMenu([("Awarded", 0, None), ("In progress", 1, None)], 1)
    rows = [b for b in menu.findChildren(QPushButton)]
    assert len(rows) == 2 and all(r.minimumHeight() >= 44 for r in rows)
    chosen = []
    menu.chosen.connect(chosen.append)
    rows[0].click()
    assert chosen == [0]


def test_update_note_shown_after_download(app, settings, monkeypatch):
    from galley.fellowships.update import UpdateResult
    from datetime import datetime, timezone
    page = make_page(app)
    page._update_done(UpdateResult(12, datetime.now(timezone.utc), ["x.json"]))
    page.refresh(MATCHES)
    assert "12 entries downloaded, 1 skipped" in texts(page.tab_stack.widget(MATCHES))
    assert "Check for Updates" in [b.text() for b in
                                   page.tab_stack.widget(MATCHES).findChildren(QPushButton)]


@pytest.mark.parametrize("level, state, date_label, phd_shown", [
    (1, 2, None, False),                         # Master's student: no PhD yet
    (2, 1, "PhD expected", False),        # PhD student: in progress
    (3, 0, "PhD awarded", False),               # Postdoc: awarded
    (4, 0, "PhD awarded", False),               # Faculty: awarded
    (0, None, None, True),                       # Not set: the user chooses
])
def test_career_stage_decides_the_phd_question(app, settings, level, state,
                                               date_label, phd_shown):
    page = make_page(app)
    page.show_tab(PROFILE)
    page.p_level.setCurrentIndex(level)
    if state is not None:
        assert page.p_phd_state.currentIndex() == state
    assert page.p_phd_state.parentWidget().isVisibleTo(page) == phd_shown
    if date_label:
        assert page.p_phd_label.text() == date_label
        assert page.p_phd_date.parentWidget().isVisibleTo(page)
    elif state == 2:
        assert not page.p_phd_date.parentWidget().isVisibleTo(page)


def test_postdoc_profile_saves_an_awarded_phd(app, settings):
    page = make_page(app)
    page.show_tab(PROFILE)
    page.p_level.setCurrentIndex(3)              # Postdoc
    page.save_profile()
    r = load_researcher(settings / "researcher.json")
    assert r.career_level == "postdoc" and r.phd_date is not None
    assert r.phd_expected is None


def test_real_entries_are_listed(app, tmp_path, monkeypatch):
    monkeypatch.setattr(submission, "user_profile_dir", lambda: tmp_path / "profiles")
    from galley.gui.fellowships_page import FellowshipsPage
    page = FellowshipsPage(today=TODAY)
    shown = texts(page.tab_stack.widget(MATCHES))
    assert "Howard Hughes Medical Institute" in shown
    assert "Show invented example entries" not in shown     # real list: no demo switch
    assert "never checked against the funder's page" not in shown  # only in details


def test_match_tiles_count_and_filter(app, settings):
    from galley.gui.app import Tile
    page = make_page(app)
    page.show_tab(MATCHES)
    tiles = {t.key: t for t in page.tab_stack.widget(MATCHES).findChildren(Tile)}
    assert set(tiles) == {"all", "eligible", "possibly eligible", "not eligible"}
    total = int(tiles["all"].text().split()[0])
    parts = sum(int(tiles[k].text().split()[0]) for k in tiles if k != "all")
    assert total == parts and tiles["all"].isChecked()
    key = next(k for k in ("not eligible", "possibly eligible", "eligible")
               if int(tiles[k].text().split()[0]))
    heading = {"not eligible": "NOT ELIGIBLE", "possibly eligible": "WORTH CHECKING",
               "eligible": "YOU'RE ELIGIBLE"}
    page._filter_status(key)
    shown = texts(page.tab_stack.widget(MATCHES))
    assert heading[key] in shown
    assert not [h for k, h in heading.items() if k != key and h in shown]


def test_application_tiles_group_the_stages(app, settings):
    from galley.gui.app import Tile
    page = make_page(app)
    page._add_application("example-international-postdoc")
    page.apps["example-international-postdoc"].set_status("interview", TODAY)
    page._add_application("example-travel-grant")
    page.refresh(APPLICATIONS)
    tiles = {t.key: t.text().split()[0]
             for t in page.tab_stack.widget(APPLICATIONS).findChildren(Tile)}
    assert tiles == {"all": "2", "preparing": "1", "review": "1",
                     "awarded": "0", "closed": "0"}
    page._filter_stage("awarded")
    assert "None at this stage" in texts(page.tab_stack.widget(APPLICATIONS))
