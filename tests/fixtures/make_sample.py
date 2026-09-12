"""Build tests/fixtures/sample_manuscript.docx: a fictional paper with known errors.

Deliberate problems (the tests expect exactly these):
  * Figure 2 is first cited before Figure 1
  * Figure 3D is cited, but the Figure 3 legend only has panels A-C
  * Figure 5 is cited but has no legend
  * Figure 4 has a legend but is never cited
  * Table 2 is cited but has no caption
  * Supplementary Figure S2 has a legend but is never cited
It also exercises the parser with:
  * Zotero, Mendeley, and EndNote citation field codes
  * a tracked insertion containing "Fig. 2B" (should count)
  * a tracked deletion containing "Fig. 6" (should be ignored)
All authors, titles, and DOIs are fictional.
"""
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

W_DATE = "2026-09-01T10:00:00Z"


def _run(text, deleted=False):
    r = OxmlElement("w:r")
    t = OxmlElement("w:delText" if deleted else "w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def _fld_char(kind):
    r = OxmlElement("w:r")
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), kind)
    r.append(fc)
    return r


def add_field(paragraph, instr, result):
    p = paragraph._p
    p.append(_fld_char("begin"))
    r = OxmlElement("w:r")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r.append(it)
    p.append(r)
    p.append(_fld_char("separate"))
    p.append(_run(result))
    p.append(_fld_char("end"))


def add_tracked(paragraph, text, kind, rid):
    wrapper = OxmlElement(f"w:{kind}")
    wrapper.set(qn("w:id"), str(rid))
    wrapper.set(qn("w:author"), "Reviewer")
    wrapper.set(qn("w:date"), W_DATE)
    wrapper.append(_run(text, deleted=(kind == "del")))
    paragraph._p.append(wrapper)


ZOTERO = ('ADDIN ZOTERO_ITEM CSL_CITATION {"citationID":"a1","citationItems":[{"id":101,'
          '"itemData":{"type":"article-journal","title":"A fictional study of gut microbes",'
          '"author":[{"family":"Doe","given":"Jane"}],"issued":{"date-parts":[[2021]]},'
          '"DOI":"10.0000/fictional.2021.001"}}],"schema":"https://github.com/'
          'citation-style-language/schema/raw/master/csl-citation.json"}')
MENDELEY = ('ADDIN CSL_CITATION {"citationID":"b2","citationItems":[{"id":"x2",'
            '"itemData":{"type":"article-journal","title":"An imaginary arthritis cohort",'
            '"author":[{"family":"Roe","given":"Sam"}],"issued":{"date-parts":[[2019]]}}}],'
            '"properties":{"noteIndex":0}}')
ENDNOTE = ('ADDIN EN.CITE <EndNote><Cite><Author>Poe</Author><Year>2020</Year>'
           '<RecNum>3</RecNum><record><titles><title>A made-up mouse model</title>'
           '</titles></record></Cite></EndNote>')


def build(path: Path):
    d = docx.Document()
    d.add_heading("Gut microbes shape arthritis severity in a fictional mouse model", 0)

    d.add_heading("Abstract", 1)
    d.add_paragraph("We tested whether fecal microbiota transplantation (FMT) changes "
                    "arthritis severity in mice.")

    d.add_heading("Introduction", 1)
    p = d.add_paragraph("Gut microbes influence immune responses ")
    add_field(p, ZOTERO, "(Doe et al., 2021)")
    p.add_run(" and joint inflammation ")
    add_field(p, MENDELEY, "(Roe et al., 2019)")
    p.add_run(". Mouse models allow causal tests ")
    add_field(p, ENDNOTE, "(Poe, 2020)")
    p.add_run(", although manual citations also appear (Moe et al., 2018).")

    d.add_heading("Results", 1)
    d.add_paragraph("Mice receiving FMT from responders showed lower arthritis scores "
                    "(Fig. 2A).")                                    # 2 before 1
    d.add_paragraph("Baseline characteristics were similar between groups "
                    "(Figure 1A\u2013C; Table 1).")
    p = d.add_paragraph("Weight loss was also reduced")
    add_tracked(p, " (Fig. 2B)", "ins", 1)                           # counts
    add_tracked(p, " (Fig. 6)", "del", 2)                            # ignored
    p.add_run(".")
    d.add_paragraph("Microbial diversity differed in Figure 3A,B and Figure 3D.")  # no 3D
    d.add_paragraph("Cytokine levels are summarized in Table 2 (Fig. 5).")         # no legends
    d.add_paragraph("Sequencing quality is shown in Supplementary Fig. S1.")

    d.add_heading("Discussion", 1)
    d.add_paragraph("Together, these data suggest that gut microbes shape arthritis "
                    "severity (Figs. 1 and 2).")

    d.add_heading("References", 1)
    for ref in ["Doe J, et al. A fictional study of gut microbes. J Fict Biol. 2021.",
                "Roe S, et al. An imaginary arthritis cohort. Fict Rheum. 2019.",
                "Poe A. A made-up mouse model. Imag Immunol. 2020.",
                "Moe B, et al. Another invented paper. Fict Med. 2018."]:
        d.add_paragraph(ref)

    d.add_heading("Tables", 1)
    d.add_paragraph("Table 1. Baseline characteristics of recipient mice.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Group", "Mean weight (g)"
    t.cell(1, 0).text, t.cell(1, 1).text = "Responder FMT", "21.4"

    d.add_heading("Figure Legends", 1)
    d.add_paragraph("Figure 1. Study design. (A) Donor selection. (B) FMT schedule. "
                    "(C) Scoring timeline.")
    d.add_paragraph("Figure 2. Arthritis outcomes. (A) Clinical scores. (B) Body weight.")
    d.add_paragraph("Figure 3. Microbial diversity. (A) Shannon index. (B) Richness. "
                    "(C) Beta diversity.")
    d.add_paragraph("Figure 4. Histology of ankle joints.")          # never cited
    d.add_paragraph("Supplementary Figure S1. Sequencing depth per sample.")
    d.add_paragraph("Supplementary Figure S2. Rarefaction curves.")  # never cited

    d.save(path)


if __name__ == "__main__":
    out = Path(__file__).with_name("sample_manuscript.docx")
    build(out)
    print(f"Wrote {out}")
