"""Build a corpus of small manuscripts, one per citation style.

Each manuscript is fictional and carries a known set of planted errors, listed
in `CORPUS`. tests/test_corpus.py checks that the tool finds exactly those and
nothing else, which is how we know the checks generalize beyond the handful of
real papers used during development.

Run directly to (re)build every file:  python tests/fixtures/make_corpus.py
"""
from __future__ import annotations

import re
import shutil
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).parent
CORPUS_DIR = HERE / "corpus"

# ---------------------------------------------------------------- helpers ---
NOTE_STYLES = {"footnote": "footnotes", "endnote": "endnotes"}


def _run(text: str, superscript: bool = False):
    r = OxmlElement("w:r")
    if superscript:
        rpr = OxmlElement("w:rPr")
        va = OxmlElement("w:vertAlign")
        va.set(qn("w:val"), "superscript")
        rpr.append(va)
        r.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def add_rich(paragraph, text: str):
    """Add text where ^{...} marks superscript and @{n} marks a note reference."""
    for part in re.split(r"(\^\{[^}]*\}|@\{\d+\})", text):
        if not part:
            continue
        if part.startswith("^{"):
            paragraph._p.append(_run(part[2:-1], superscript=True))
        elif part.startswith("@{"):
            r = OxmlElement("w:r")
            ref = OxmlElement("w:footnoteReference")
            ref.set(qn("w:id"), part[2:-1])
            r.append(ref)
            paragraph._p.append(r)
        else:
            paragraph._p.append(_run(part))


FOOTNOTES_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>
<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/>
</w:r></w:p></w:footnote>
{notes}
</w:footnotes>"""
NOTE_XML = ('<w:footnote w:id="{id}"><w:p><w:r><w:t xml:space="preserve">{text}</w:t>'
            '</w:r></w:p></w:footnote>')
REL = ('<Relationship Id="rIdNotes" Type="http://schemas.openxmlformats.org/'
       'officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/>')
CT = ('<Override PartName="/word/footnotes.xml" ContentType="application/'
      'vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>')


def _escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def inject_footnotes(path: Path, notes: dict[int, str]):
    """Add a footnotes part to a saved .docx (python-docx can't create one)."""
    temp = path.with_suffix(".tmp.docx")
    body = "\n".join(NOTE_XML.format(id=i, text=_escape(t)) for i, t in sorted(notes.items()))
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(
            temp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", CT.encode() + b"</Types>")
            elif item.filename == "word/_rels/document.xml.rels":
                data = data.replace(b"</Relationships>", REL.encode() + b"</Relationships>")
            dst.writestr(item, data)
        dst.writestr("word/footnotes.xml",
                     FOOTNOTES_XML.format(notes=body))
    shutil.move(temp, path)


# ------------------------------------------------------------ corpus spec ---
@dataclass
class Manuscript:
    name: str
    style: str                      # how citations are written, for the docstring
    body: list[str]                 # Results paragraphs (rich text markup allowed)
    references: list[str]
    legends: list[str] = field(default_factory=list)
    notes: dict[int, str] = field(default_factory=dict)
    numbered_list: bool = False     # reference list uses Word auto-numbering
    expected: list[str] = field(default_factory=list)   # substrings of expected messages
    expect_clean: bool = False      # no errors or warnings at all


REFS_NUMBERED = [
    "1\tAdler K, Byrne T. Gut communities in model hosts. J Fict Biol 2019;4:1-12.",
    "2\tCole R, Danver P. Microbial metabolism of small molecules. Fict Chem 2020;7:55-70.",
    "3\tEverly S. Host immunity and diet. Imag Immunol 2021;12:300-15.",
    "4\tFontaine L, Guo M. Sequencing depth and diversity estimates. Fict Methods 2018;2:9-20.",
]
REFS_PLAIN = [
    "Adler K, Byrne T. Gut communities in model hosts. J Fict Biol 2019;4:1-12.",
    "Cole R, Danver P. Microbial metabolism of small molecules. Fict Chem 2020;7:55-70.",
    "Everly S. Host immunity and diet. Imag Immunol 2021;12:300-15.",
]
REFS_AUTHOR_YEAR = [
    "Adler, K., & Byrne, T. (2019). Gut communities in model hosts. Journal of Fictional "
    "Biology, 4, 1-12.",
    "Cole, R., & Danver, P. (2020). Microbial metabolism of small molecules. Fictional "
    "Chemistry, 7, 55-70.",
    "Everly, S. (2021). Host immunity and diet. Imaginary Immunology, 12, 300-315.",
]
LEGENDS_2 = ["Figure 1. Study design. (A) Groups. (B) Timeline.",
             "Figure 2. Community composition. (A) Genera. (B) Diversity."]

CORPUS = [
    Manuscript(
        name="vancouver_brackets", style="numbered, [1]",
        body=["Communities differed between groups [1] (Figure 1A).",
              "Metabolism varied by strain [2,3] (Figure 2A).",
              "Depth was sufficient throughout [4] (Figure 2B)."],
        references=REFS_NUMBERED, legends=LEGENDS_2, expect_clean=True),
    Manuscript(
        name="vancouver_uncited", style="numbered, [1], with a planted error",
        body=["Communities differed between groups [1] (Figure 1A).",
              "Metabolism varied by strain [2] (Figure 2A).",
              "Depth was sufficient [4] (Figure 2B)."],
        references=REFS_NUMBERED, legends=LEGENDS_2,
        expected=["Reference 3 is never cited in the text."]),
    Manuscript(
        name="vancouver_overrun", style="numbered, citation beyond the list",
        body=["Communities differed [1] (Figure 1A).",
              "Metabolism varied [2,3] (Figure 2A).",
              "Later work disagrees [9] (Figure 2B).",
              "Depth was sufficient [4]."],
        references=REFS_NUMBERED, legends=LEGENDS_2,
        expected=["Citation 9 has no matching reference (the list has 4 entries)."]),
    Manuscript(
        name="nature_superscript", style="superscript numbers",
        body=["Communities differed between groups^{1} (Figure 1A).",
              "Metabolism varied by strain^{2,3} (Figure 2A).",
              "Depth was sufficient throughout^{4} (Figure 2B)."],
        references=REFS_NUMBERED, legends=LEGENDS_2, expect_clean=True),
    Manuscript(
        name="word_autonumbered_list", style="numbered list with no typed numbers",
        body=["Communities differed between groups [1] (Figure 1A).",
              "Metabolism varied by strain [2] (Figure 2A).",
              "Diet mattered as well [3] (Figure 2B)."],
        references=REFS_PLAIN, legends=LEGENDS_2, numbered_list=True, expect_clean=True),
    Manuscript(
        name="apa_author_year", style="(Adler & Byrne, 2019)",
        body=["Communities differed between groups (Adler & Byrne, 2019); see Figure 1A.",
              "Metabolism varied by strain (Cole & Danver, 2020), as in Figure 2A.",
              "Everly (2021) reported a similar effect (Figure 2B)."],
        references=REFS_AUTHOR_YEAR, legends=LEGENDS_2, expect_clean=True),
    Manuscript(
        name="harvard_missing_entry", style="author-year with a citation that isn't listed",
        body=["Communities differed between groups (Adler and Byrne, 2019); see Figure 1A.",
              "An earlier survey disagreed (Mortimer et al., 2015).",
              "Metabolism varied by strain (Cole and Danver, 2020), as in Figure 2A.",
              "Diet mattered too (Everly, 2021) (Figure 2B)."],
        references=REFS_AUTHOR_YEAR, legends=LEGENDS_2,
        expected=["The citation to Mortimer (2015) has no matching entry"]),
    Manuscript(
        name="bracket_author_year", style="[Adler and Byrne 2019]",
        body=["Communities differed between groups [Adler and Byrne 2019] (Figure 1A).",
              "Metabolism varied by strain [Cole and Danver 2020] (Figure 2A).",
              "Diet mattered too [Everly 2021] (Figure 2B)."],
        references=REFS_AUTHOR_YEAR, legends=LEGENDS_2, expect_clean=True),
    Manuscript(
        name="latex_keys", style="[Adl19], [CD20]",
        body=["Communities differed between groups [Adl19] (Figure 1A).",
              "Metabolism varied by strain [CD20] (Figure 2A).",
              "Diet mattered too [Eve21] (Figure 2B)."],
        references=REFS_AUTHOR_YEAR, legends=LEGENDS_2, expect_clean=True),
    Manuscript(
        name="latex_keys_missing", style="[Adl19] with a key that isn't listed",
        body=["Communities differed between groups [Adl19] (Figure 1A).",
              "An earlier survey disagreed [Mor15].",
              "Metabolism varied by strain [CD20] (Figure 2A).",
              "Diet mattered too [Eve21] (Figure 2B)."],
        references=REFS_AUTHOR_YEAR, legends=LEGENDS_2,
        expected=["[Mor15] has no matching entry"]),
    Manuscript(
        name="chicago_footnotes", style="footnote citations",
        body=["Communities differed between groups@{1} (Figure 1A).",
              "Metabolism varied by strain@{2} (Figure 2A).",
              "Diet mattered too@{3} (Figure 2B)."],
        references=REFS_AUTHOR_YEAR, legends=LEGENDS_2,
        notes={1: "Adler and Byrne, \u201cGut communities in model hosts,\u201d 2019, 1-12.",
               2: "Cole and Danver, \u201cMicrobial metabolism,\u201d 2020, 55-70.",
               3: "Everly, \u201cHost immunity and diet,\u201d 2021, 300-315."},
        expect_clean=True),
    Manuscript(
        name="footnotes_uncited_entry", style="footnote citations, one entry unused",
        body=["Communities differed between groups@{1} (Figure 1A).",
              "Metabolism varied by strain@{2} (Figure 2A and 2B)."],
        references=REFS_AUTHOR_YEAR, legends=LEGENDS_2,
        notes={1: "Adler and Byrne, \u201cGut communities in model hosts,\u201d 2019, 1-12.",
               2: "Cole and Danver, \u201cMicrobial metabolism,\u201d 2020, 55-70."},
        expected=["never cited in the text"]),
    Manuscript(
        name="no_citations_at_all", style="a draft with no citations yet",
        body=["Communities differed between groups (Figure 1A).",
              "Metabolism varied by strain (Figure 2A and 2B)."],
        references=[], legends=LEGENDS_2, expect_clean=True),
    Manuscript(
        name="unrecognized_style", style="an in-house numbering scheme",
        body=["Communities differed between groups {ref: AB-19} (Figure 1A).",
              "Metabolism varied by strain {ref: CD-20} (Figure 2A and 2B)."],
        references=REFS_AUTHOR_YEAR, legends=LEGENDS_2,
        expected=["weren't compared"]),
]


def build(spec: Manuscript, out_dir: Path = CORPUS_DIR) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{spec.name}.docx"
    d = docx.Document()
    d.add_heading("A fictional study of host-associated microbial communities", 0)
    d.add_heading("Results", 1)
    for text in spec.body:
        add_rich(d.add_paragraph(), text)
    if spec.references:
        d.add_heading("References", 1)
        for ref in spec.references:
            p = d.add_paragraph(ref, style="List Number" if spec.numbered_list else None)
            if spec.numbered_list:
                p.style = d.styles["List Number"]
    if spec.legends:
        d.add_heading("Figure Legends", 1)
        for legend in spec.legends:
            d.add_paragraph(legend)
    d.save(path)
    if spec.notes:
        inject_footnotes(path, spec.notes)
    return path


def build_all(out_dir: Path = CORPUS_DIR) -> list[Path]:
    return [build(spec, out_dir) for spec in CORPUS]


if __name__ == "__main__":
    for p in build_all():
        print("Wrote", p)
