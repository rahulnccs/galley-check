"""Application requirements, preparation plans, the tracker and reminders."""
import zipfile
from datetime import date, timedelta
from urllib.parse import parse_qs, urlparse

import docx
import pytest

from galley.fellowships import (Fellowship, check_application, check_document,
                                load_applications, plan, reminders,
                                save_applications, start_application)
from galley.fellowships.application_check import (FAIL, PASS, WARN,
                                                  _docx_saved_pages)
from galley.fellowships.requirements import Requirements

TODAY = date(2026, 10, 1)
FINAL = date(2027, 1, 15)


def fellowship(**overrides) -> Fellowship:
    data = {"id": "f", "name": "Test Fellowship", "funder": "Test Funder",
            "url": "https://example.org/f", "verified": "2026-09-01",
            "deadlines": [{"kind": "final", "date": FINAL.isoformat()}]}
    data.update(overrides)
    return Fellowship.from_dict(data)


FULL = {"documents": [{"name": "Research proposal", "max_words": 50,
                       "sections": ["Background", "Aims"]}],
        "host_letter": True}


# --- requirements ------------------------------------------------------------

@pytest.mark.parametrize("bad, complaint", [
    ({"documents": [{"max_pages": 3}]}, "needs a name"),
    ({"cv_format": "fancy"}, "cv_format must be"),
    ({"documents": [{"name": "Proposal", "max_words": "lots"}]}, "invalid literal"),
    ({"submission": "fax"}, "submission must be"),
])
def test_bad_requirements_rejected(bad, complaint):
    with pytest.raises(ValueError, match=complaint):
        fellowship(requirements=bad)


def test_requirements_parsed():
    f = fellowship(requirements=FULL)
    assert f.requirements.host_letter
    assert f.requirements.document("research proposal").max_words == 50


def make_draft(path, paragraphs, pages=None, size=None, margin_cm=None):
    from docx.shared import Cm, Pt
    d = docx.Document()
    for kind, text in paragraphs:
        p = d.add_heading(text, level=1) if kind == "h" else d.add_paragraph(text)
        if size and kind != "h":
            for run in p.runs:
                run.font.size = Pt(size)
    if margin_cm is not None:
        for section in d.sections:
            section.left_margin = section.right_margin = Cm(margin_cm)
            section.top_margin = section.bottom_margin = Cm(margin_cm)
    d.save(path)
    if pages is not None:   # what Word writes on save
        tmp = str(path) + ".tmp"
        with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w") as dst:
            for item in src.infolist():
                if item.filename != "docProps/app.xml":
                    dst.writestr(item, src.read(item.filename))
            dst.writestr("docProps/app.xml",
                         f"<Properties><Pages>{pages}</Pages></Properties>")
        import os
        os.replace(tmp, path)
    return str(path)


def outcomes(findings):
    return [(f.outcome, f.text) for f in findings]


def doc_req(**kw):
    return Requirements.from_dict({"documents": [dict(name="Research proposal", **kw)]}
                                  ).documents[0]


def test_draft_within_limits(tmp_path):
    path = make_draft(tmp_path / "p.docx", [("h", "Background"), ("p", "Short text."),
                                            ("h", "Aims"), ("p", "Two aims.")],
                      pages=3, size=12, margin_cm=2.5)
    req = doc_req(max_words=50, max_pages=5, min_font_size=11, min_margin_cm=2,
                  sections=["Background", "Aims"])
    found = outcomes(check_document(path, req))
    assert not [t for o, t in found if o != PASS], found
    assert (PASS, "3 of 5 pages (as last saved in Word).") in found


def test_draft_over_word_limit_and_missing_section(tmp_path):
    path = make_draft(tmp_path / "p.docx",
                      [("h", "Background"), ("p", "word " * 60)])
    found = outcomes(check_document(path, doc_req(max_words=50,
                                                  sections=["Background", "Aims"])))
    assert (FAIL, "60 words, over the 50-word limit by 10.") in found
    assert (FAIL, 'No "Aims" section was found.') in found
    assert (PASS, 'Includes "Background".') in found


def test_page_count_from_word(tmp_path):
    over = make_draft(tmp_path / "over.docx", [("p", "x")], pages=7)
    found = outcomes(check_document(over, doc_req(max_pages=5)))
    assert any(o == FAIL and "7 pages" in t for o, t in found)
    unknown = make_draft(tmp_path / "unknown.docx", [("p", "x")])
    if _docx_saved_pages(unknown) is None:
        found = outcomes(check_document(unknown, doc_req(max_pages=5)))
        assert any(o == WARN and "can't count pages" in t for o, t in found)


def test_small_text_fails(tmp_path):
    path = make_draft(tmp_path / "p.docx", [("p", "word " * 40)], size=9)
    found = outcomes(check_document(path, doc_req(min_font_size=11)))
    assert any(o == FAIL and "smaller than the 11 pt minimum" in t and "9 pt" in t
               for o, t in found)


def test_narrow_margins_fail(tmp_path):
    path = make_draft(tmp_path / "p.docx", [("p", "text")], margin_cm=1.27)
    found = outcomes(check_document(path, doc_req(min_margin_cm=2)))
    assert any(o == FAIL and "narrower than 2 cm" in t and "left 1.3 cm" in t
               for o, t in found)


def test_word_file_when_pdf_wanted(tmp_path):
    path = make_draft(tmp_path / "p.docx", [("p", "text")])
    found = outcomes(check_document(path, doc_req(file_format="pdf")))
    assert [(o, t) for o, t in found if "as a PDF" in t] == [
        ("info", "Submit the Research proposal as a PDF: export it from Word "
                 "when it's final, and check the PDF's page count.")]


def test_bad_file_format_rejected():
    with pytest.raises(ValueError, match="file_format"):
        doc_req(file_format="odt")


def test_unsupported_file_type(tmp_path):
    other = tmp_path / "p.txt"
    other.write_text("hello")
    with pytest.raises(ValueError, match="isn't a Word document or a PDF"):
        check_document(str(other), doc_req())


def test_whole_application(tmp_path):
    f = fellowship(requirements={
        "documents": [{"name": "Research proposal", "max_words": 50},
                      {"name": "Career statement", "max_words": 20}],
        "cv_format": "narrative", "host_letter": True})
    proposal = make_draft(tmp_path / "p.docx", [("p", "Fine and short.")])
    report = check_application(f, {"Research proposal": proposal})
    by_name = {d.name: d for d in report.documents}
    assert by_name["Research proposal"].problems == 0
    assert outcomes(by_name["Career statement"].findings)[0] == (FAIL, "Not attached yet.")
    assert not report.ready and report.problems == 1
    assert any("narrative CV" in g.text for g in report.general)
    assert any("letter of support" in g.text for g in report.general)

    statement = make_draft(tmp_path / "s.docx", [("p", "Brief.")])
    report = check_application(f, {"Research proposal": proposal,
                                   "Career statement": statement})
    assert report.ready


def test_moved_file_is_reported(tmp_path):
    f = fellowship(requirements={"documents": [{"name": "Proposal"}]})
    report = check_application(f, {"Proposal": str(tmp_path / "gone.docx")})
    assert "can no longer be found" in report.documents[0].findings[0].text


# --- preparation plan --------------------------------------------------------

def test_plan_works_back_from_the_deadline():
    f = fellowship(requirements=FULL)
    steps = {s.key: s for s in plan(f, TODAY)}
    assert steps["submit"].date == FINAL and steps["submit"].is_deadline
    assert steps["drafts_for_feedback"].date == FINAL - timedelta(weeks=4)
    assert steps["contact_hosts"].date == FINAL - timedelta(weeks=12)
    assert steps["final_check"].date == FINAL - timedelta(weeks=1)
    assert "Research proposal" in steps["start_writing"].task
    dates = [s.date for s in plan(f, TODAY)]
    assert dates == sorted(dates)


def test_internal_deadline_pulls_drafting_earlier():
    internal = FINAL - timedelta(weeks=4)
    f = fellowship(deadlines=[{"kind": "internal", "date": internal.isoformat()},
                              {"kind": "final", "date": FINAL.isoformat()}])
    steps = {s.key: s for s in plan(f, TODAY)}
    assert steps["internal"].is_deadline
    assert steps["start_writing"].date == internal - timedelta(weeks=8)


def test_no_plan_without_a_deadline():
    assert plan(fellowship(deadlines=[], rolling=True), TODAY) == []


def test_late_start_keeps_steps_as_overdue():
    late = FINAL - timedelta(weeks=2)
    steps = {s.key: s for s in plan(fellowship(requirements=FULL), late)}
    assert steps["contact_hosts"].overdue(late)
    assert not steps["submit"].overdue(late)


# --- tracker -----------------------------------------------------------------

def test_tracker_round_trip(tmp_path):
    path = tmp_path / "applications.json"
    apps = {}
    app = start_application(apps, fellowship(), TODAY)
    app.mark_done("start_writing")
    app.set_status("submitted", TODAY)
    app.notes = "Host: Dr Example"
    save_applications(apps, path)
    back = load_applications(path)["f"]
    assert back.status == "submitted" and back.done == ["start_writing"]
    assert back.status_dates["submitted"] == TODAY.isoformat()
    assert back.notes == "Host: Dr Example"


def test_starting_twice_keeps_the_original(tmp_path):
    apps = {}
    a = start_application(apps, fellowship(), TODAY)
    a.notes = "keep"
    assert start_application(apps, fellowship(), TODAY).notes == "keep"


def test_bad_status_rejected():
    app = start_application({}, fellowship(), TODAY)
    with pytest.raises(ValueError):
        app.set_status("pending", TODAY)


def test_missing_or_broken_file_gives_no_applications(tmp_path):
    assert load_applications(tmp_path / "none.json") == {}
    (tmp_path / "bad.json").write_text("{ broken")
    assert load_applications(tmp_path / "bad.json") == {}


# --- reminders ---------------------------------------------------------------

def test_reminders_when_the_app_opens():
    f = fellowship(requirements=FULL)
    apps = {}
    start_application(apps, f, TODAY)
    day = FINAL - timedelta(weeks=5)
    found = reminders(apps, [f], day)
    texts = [r.text for r in found]
    assert any("Contact potential host labs" in t for t in texts)
    assert found[0].urgency == "overdue"            # most urgent first
    assert all((r.date - day).days <= 42 for r in found)


def test_done_and_snoozed_steps_are_quiet():
    f = fellowship(requirements=FULL)
    apps = {}
    app = start_application(apps, f, TODAY)
    day = FINAL - timedelta(weeks=5)
    app.mark_done("contact_hosts")
    app.snooze("drafts_for_feedback", day + timedelta(days=7))
    keys = [r.step for r in reminders(apps, [f], day)]
    assert "contact_hosts" not in keys and "drafts_for_feedback" not in keys


def test_submitted_application_stops_preparation_reminders():
    f = fellowship(requirements=FULL)
    apps = {}
    start_application(apps, f, TODAY).set_status("submitted", TODAY)
    assert reminders(apps, [f], FINAL - timedelta(weeks=1)) == []


def test_interview_reminder():
    f = fellowship()
    apps = {}
    app = start_application(apps, f, TODAY)
    app.set_status("interview", TODAY)
    app.interview_date = (TODAY + timedelta(days=10)).isoformat()
    found = reminders(apps, [f], TODAY)
    assert [(r.text, r.urgency) for r in found] == [("Interview", "soon")]


def test_moved_deadline_reported_once():
    apps = {}
    start_application(apps, fellowship(), TODAY)
    moved = fellowship(deadlines=[{"kind": "final", "date": "2027-02-01"}])
    first = reminders(apps, [moved], TODAY)
    assert any(r.urgency == "changed" and "15 Jan 2027 to 01 Feb 2027" in r.text
               for r in first)
    assert not [r for r in reminders(apps, [moved], TODAY) if r.urgency == "changed"]


# --- reporting a problem with an entry ----------------------------------------

def test_report_link_names_the_entry_only():
    url = fellowship().report_problem_url()
    q = parse_qs(urlparse(url).query)
    assert url.startswith("https://github.com/rahulnccs/galley-check/issues/new")
    assert "Test Fellowship" in q["title"][0]
    assert "https://example.org/f" in q["body"][0]


def test_no_report_link_for_your_own_entries():
    f = Fellowship.from_dict({"id": "custom-x", "name": "X", "custom": True})
    assert f.report_problem_url() is None


def test_pdf_pages_text_size_and_margins(tmp_path):
    pytest.importorskip("pdfplumber")
    pytest.importorskip("PySide6")
    from PySide6.QtCore import QMarginsF, QRectF, Qt
    from PySide6.QtGui import QFont, QPageLayout, QPageSize, QPainter, QPdfWriter
    from PySide6.QtWidgets import QApplication
    if QApplication.instance() is None:
        try:
            QApplication([])
        except Exception as e:
            pytest.skip(f"Qt cannot start here: {e}")

    def make(path, pt, margin_mm, pages):
        w = QPdfWriter(str(path))
        w.setPageSize(QPageSize(QPageSize.A4))
        w.setResolution(72)
        w.setPageMargins(QMarginsF(*[margin_mm] * 4), QPageLayout.Millimeter)
        p = QPainter(w)
        font = QFont()
        font.setPointSizeF(pt)
        p.setFont(font)
        for i in range(pages):
            if i:
                w.newPage()
            p.drawText(QRectF(0, 0, w.width(), 40), Qt.AlignLeft, "Background")
            p.drawText(QRectF(0, 60, w.width(), 400), Qt.TextWordWrap,
                       "Aims of the project. " + "word " * 120)
        p.end()
        return str(path)

    req = doc_req(max_pages=3, min_font_size=11, min_margin_cm=2,
                  sections=["Background", "Aims"], file_format="pdf")
    good = outcomes(check_document(make(tmp_path / "ok.pdf", 12, 25, 2), req))
    assert all(o == PASS for o, _ in good), good
    bad = outcomes(check_document(make(tmp_path / "bad.pdf", 9, 10, 4), req))
    assert (FAIL, "4 pages, over the 3-page limit.") in bad
    assert any(o == FAIL and "9 pt" in t for o, t in bad)
    assert any(o == FAIL and "Margins are narrower" in t for o, t in bad)
