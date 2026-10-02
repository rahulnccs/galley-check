"""Check a fellowship application's documents against the funder's format.

For each required document the user attaches a file (.docx, or .pdf when
pdfplumber is installed), and Galley checks what it can measure reliably:

  * page count (exact for PDF; the count Word saved, for .docx)
  * word count
  * required sections, by heading or run-in heading
  * smallest text size
  * page margins (all four from a .docx's page setup; left and right for
    a PDF, where headers and footers make top and bottom unreliable)
  * the file format the funder asks for

Every finding says whether it passed, failed, or needs the user to look, so
the report doubles as a checklist. Nothing is sent anywhere.
"""
from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from .model import Fellowship
from .requirements import DocumentRequirement

PASS, FAIL, WARN, INFO = "pass", "fail", "warn", "info"
WORD = re.compile(r"[\wÀ-ɏ][\wÀ-ɏ'’\-]*")
PT_PER_CM = 72 / 2.54
# Small text that is allowed to be small: superscript citation numbers,
# footnote markers, page numbers. Up to this share of characters is fine.
SMALL_TEXT_TOLERANCE = 0.02


@dataclass
class Finding:
    outcome: str            # PASS | FAIL | WARN | INFO
    text: str
    suggestion: str | None = None


@dataclass
class DocumentReport:
    name: str
    path: str | None
    findings: list[Finding] = field(default_factory=list)

    @property
    def problems(self) -> int:
        return sum(1 for f in self.findings if f.outcome == FAIL)

    @property
    def warnings(self) -> int:
        return sum(1 for f in self.findings if f.outcome == WARN)


@dataclass
class ApplicationReport:
    fellowship: Fellowship
    documents: list[DocumentReport] = field(default_factory=list)
    general: list[Finding] = field(default_factory=list)

    @property
    def problems(self) -> int:
        return (sum(d.problems for d in self.documents)
                + sum(1 for f in self.general if f.outcome == FAIL))

    @property
    def warnings(self) -> int:
        return (sum(d.warnings for d in self.documents)
                + sum(1 for f in self.general if f.outcome == WARN))

    @property
    def ready(self) -> bool:
        return self.problems == 0


# ---- reading files ---------------------------------------------------------

@dataclass
class Measured:
    """What could be measured in one file. None means it couldn't be."""
    kind: str                           # "docx" | "pdf"
    words: int = 0
    pages: int | None = None
    pages_exact: bool = False
    headings: list[str] = field(default_factory=list)   # candidate headings
    sizes: list[tuple[float, int]] = field(default_factory=list)  # (pt, chars)
    margins_cm: dict[str, float] = field(default_factory=dict)


def pdf_supported() -> bool:
    try:
        import pdfplumber  # noqa: F401
    except ImportError:
        return False
    return True


def _docx_saved_pages(path: str) -> int | None:
    """The page count Word stored on last save, in docProps/app.xml. Other
    editors may not write it, and it is stale if the file was edited since."""
    try:
        with zipfile.ZipFile(path) as z:
            xml = z.read("docProps/app.xml").decode("utf-8", "replace")
    except (KeyError, OSError, zipfile.BadZipFile):
        return None
    m = re.search(r"<(?:\w+:)?Pages>(\d+)</", xml)
    return int(m.group(1)) if m else None


def _docx_default_size(d) -> float | None:
    """The document's default text size (docDefaults), in points."""
    from docx.oxml.ns import qn
    styles = d.styles.element
    sz = styles.find(f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/"
                     f"{qn('w:rPr')}/{qn('w:sz')}")
    if sz is not None and sz.get(qn("w:val")):
        return int(sz.get(qn("w:val"))) / 2
    return None


def _style_size(style) -> float | None:
    while style is not None:
        if style.font is not None and style.font.size is not None:
            return style.font.size.pt
        style = style.base_style
    return None


def _measure_docx(path: str) -> Measured:
    import docx

    from ..engine import load

    m = Measured("docx")
    parsed = load(path)
    body = [p for p in parsed.paragraphs if p.text]
    m.words = sum(len(WORD.findall(p.text)) for p in body if not p.is_heading)
    m.headings = ([p.text.lower() for p in body if p.is_heading]
                  + [p.text[:70].lower() for p in body if not p.is_heading])
    m.pages = _docx_saved_pages(path)

    d = docx.Document(path)
    default = _docx_default_size(d) or 10.0     # Word's own default is 10 pt
    sizes = m.sizes
    for para in d.paragraphs:
        para_size = _style_size(para.style)
        for run in para.runs:
            text = run.text.strip()
            if not text:
                continue
            if run.font.superscript or run.font.subscript:
                continue                         # citation markers, formulas
            size = (run.font.size.pt if run.font.size is not None
                    else _style_size(run.style) if run.style is not None
                    and run.style.name != "Default Paragraph Font" else None)
            sizes.append((size or para_size or default, len(text)))
    m.margins_cm = {}
    for section in d.sections:
        for side in ("left", "right", "top", "bottom"):
            value = getattr(section, f"{side}_margin")
            if value is not None:
                cm = value.cm
                m.margins_cm[side] = min(cm, m.margins_cm.get(side, cm))
    return m


def _measure_pdf(path: str) -> Measured:
    import pdfplumber

    m = Measured("pdf", pages_exact=True)
    sizes = m.sizes
    left = right = None
    with pdfplumber.open(path) as pdf:
        m.pages = len(pdf.pages)
        for page in pdf.pages:
            text = page.extract_text() or ""
            m.words += len(WORD.findall(text))
            m.headings += [line[:70].lower() for line in text.splitlines() if line.strip()]
            chars = [c for c in page.chars if c.get("text", "").strip()]
            for c in chars:
                sizes.append((round(float(c["size"]), 1), 1))
            if chars:
                x0 = min(c["x0"] for c in chars) / PT_PER_CM
                x1 = (float(page.width) - max(c["x1"] for c in chars)) / PT_PER_CM
                left = x0 if left is None else min(left, x0)
                right = x1 if right is None else min(right, x1)
    if left is not None:
        m.margins_cm = {"left": left, "right": right}
    return m


def measure(path: str) -> Measured:
    suffix = Path(path).suffix.lower()
    if suffix == ".docx":
        return _measure_docx(path)
    if suffix == ".pdf":
        if not pdf_supported():
            raise ValueError("Checking PDFs needs the pdfplumber package. "
                             "Attach the Word version instead.")
        return _measure_pdf(path)
    raise ValueError(f"{Path(path).name} isn't a Word document or a PDF.")


# ---- checking --------------------------------------------------------------

def _covers(candidate: str, words: list[str]) -> bool:
    at = 0
    for w in words:
        found = candidate.find(w, at)
        if found == -1:
            return False
        at = found + len(w)
    return True


def _size_findings(m: Measured, minimum: float) -> list[Finding]:
    sizes = m.sizes
    total = sum(n for _, n in sizes)
    if not total:
        return []
    small = [(s, n) for s, n in sizes if s < minimum - 0.25]
    share = sum(n for _, n in small) / total
    smallest = min(s for s, _ in small) if small else min(s for s, _ in sizes)
    if not small:
        return [Finding(PASS, f"Text is {minimum:g} pt or larger "
                              f"(smallest {smallest:g} pt).")]
    if share <= SMALL_TEXT_TOLERANCE:
        return [Finding(WARN, f"A little text is smaller than {minimum:g} pt "
                              f"(smallest {smallest:g} pt), probably footnotes "
                              f"or labels.",
                        "Check whether the funder's minimum applies to these.")]
    return [Finding(FAIL, f"{share:.0%} of the text is smaller than the "
                          f"{minimum:g} pt minimum (smallest {smallest:g} pt).",
                    f"Set the body text to at least {minimum:g} pt.")]


def check_document(path: str, req: DocumentRequirement) -> list[Finding]:
    """Check one file against one document's requirements."""
    m = measure(path)
    out: list[Finding] = []
    name = req.name

    if req.file_format and req.file_format != m.kind:
        wanted = "PDF" if req.file_format == "pdf" else "Word document"
        if req.file_format == "pdf":
            out.append(Finding(WARN, f"The funder wants the {name} as a PDF; "
                                     f"you attached the Word file.",
                               "Export it as PDF before submitting, and check "
                               "the PDF's page count."))
        else:
            out.append(Finding(FAIL, f"The funder wants the {name} as a {wanted}.",
                               "Attach and submit the Word file."))

    if req.max_pages is not None:
        if m.pages is None:
            out.append(Finding(WARN, f"The limit is {req.max_pages} pages, but "
                                     f"Galley can't count pages in this file.",
                               "Open it in Word, save, and check again, or "
                               "attach the PDF."))
        elif m.pages > req.max_pages:
            how = "" if m.pages_exact else " when last saved in Word"
            out.append(Finding(FAIL, f"{m.pages} pages{how}, over the "
                                     f"{req.max_pages}-page limit.",
                               "Shorten it; reference lists often count too."))
        else:
            how = "" if m.pages_exact else " (as last saved in Word)"
            out.append(Finding(PASS, f"{m.pages} of {req.max_pages} pages{how}."))

    if req.max_words is not None:
        if m.words > req.max_words:
            out.append(Finding(FAIL, f"{m.words:,} words, over the "
                                     f"{req.max_words:,}-word limit by "
                                     f"{m.words - req.max_words:,}.",
                               "Check whether references count towards the "
                               "limit; Galley counts them."))
        else:
            out.append(Finding(PASS, f"{m.words:,} of {req.max_words:,} words."))

    for section in req.sections:
        wanted = [w for w in WORD.findall(section.lower())
                  if w not in {"and", "of", "the"}]
        if any(_covers(c, wanted) for c in m.headings):
            out.append(Finding(PASS, f'Includes "{section}".'))
        else:
            out.append(Finding(FAIL, f'No "{section}" section was found.',
                               "Add it, with a heading that matches the "
                               "funder's wording."))

    if req.min_font_size:
        out += _size_findings(m, req.min_font_size)

    if req.min_margin_cm:
        narrow = {side: cm for side, cm in m.margins_cm.items()
                  if cm < req.min_margin_cm - 0.05}
        if narrow:
            sides = ", ".join(f"{side} {cm:.1f} cm" for side, cm in narrow.items())
            out.append(Finding(FAIL, f"Margins are narrower than "
                                     f"{req.min_margin_cm:g} cm: {sides}.",
                               "Widen them in Layout › Margins."))
        elif m.margins_cm:
            checked = ("all sides" if len(m.margins_cm) == 4
                       else "left and right")
            out.append(Finding(PASS, f"Margins are at least "
                                     f"{req.min_margin_cm:g} cm ({checked})."))
    return out


def check_application(f: Fellowship, files: dict[str, str]) -> ApplicationReport:
    """Check every attached document. `files` maps a required document's name
    to the file the user attached for it."""
    report = ApplicationReport(f)
    req = f.requirements
    if req is None:
        report.general.append(Finding(INFO, "This fellowship lists no document "
                                            "requirements to check against."))
        return report

    for doc in req.documents:
        path = files.get(doc.name)
        dr = DocumentReport(doc.name, path)
        if not path:
            dr.findings.append(Finding(FAIL, "Not attached yet.",
                                       "Attach the file to check it."))
        elif not Path(path).exists():
            dr.findings.append(Finding(FAIL, f"{Path(path).name} can no longer be "
                                             f"found; it may have been moved.",
                                       "Attach it again."))
        else:
            try:
                dr.findings += check_document(path, doc)
            except Exception as e:      # unreadable or unsupported file
                dr.findings.append(Finding(FAIL, f"Couldn't read "
                                                 f"{Path(path).name}: {e}"))
        report.documents.append(dr)

    if req.cv_format == "narrative":
        report.general.append(Finding(INFO, "The CV must be a narrative CV."
                                      + (f" {req.cv_notes}" if req.cv_notes else "")))
    elif req.cv_format == "funder_template":
        report.general.append(Finding(INFO, "The CV must use the funder's "
                                            "template."
                                      + (f" {req.cv_notes}" if req.cv_notes else "")))
    if req.host_letter:
        report.general.append(Finding(INFO, "Remember the letter of support "
                                            "from your host."))
    return report
